"""Audit only public G6 candidate metadata and targeted-panel coverage.

Run with the existing project Python environment::

    python src/101_audit_g6_avey_dharshini_metadata.py

This script never opens a spatial count matrix or estimates an AD-control effect.
"""

from __future__ import annotations

import csv
import gzip
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/G6_source_metadata_2026-09-29"
REG = ROOT / "results/2026-09-24-disease-side-v3.5-g6-input-audit/registries"
OUT = ROOT / "results/2026-09-29-g6-avey-dharshini-metadata-audit"
AVEY_TABLE = RAW / "avey_supplementary_file5.xlsx"
AVEY_CONFOUNDS = RAW / "avey_author_confounds.csv"
AVEY_DIAGNOSIS = RAW / "avey_author_diagnosis.csv"
DHARSHINI_TABLE = RAW / "dharshini_xenium_metadata.xlsx"
BASE_PANEL = RAW / "xenium_human_brain_266.csv"
G1 = REG / "G1_microglia_signed_sender_244.tsv.gz"
G5 = REG / "G5_receiver_one_to_one_human_ortholog_registry.tsv.gz"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_tsv(path: Path, rows: list[dict], columns: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, delimiter="\t", fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def read_avey() -> tuple[list[dict], dict]:
    with AVEY_CONFOUNDS.open(newline="", encoding="utf-8-sig") as handle:
        confounds = {int(row["SampleID"].removeprefix("Sample")): row for row in csv.DictReader(handle)}
    with AVEY_DIAGNOSIS.open(newline="", encoding="utf-8-sig") as handle:
        diagnoses = {int(row["SampleID"].removeprefix("Sample")): row for row in csv.DictReader(handle)}
    sheet = load_workbook(AVEY_TABLE, read_only=True, data_only=True)["ST cases"]
    rows = list(sheet.values)
    fields = rows[0][:11]
    donors = []
    for values in rows[1:]:
        if not isinstance(values[0], int):
            continue
        source = dict(zip(fields, values[:11]))
        donor_id = int(source["ID"])
        assert donor_id in confounds and donor_id in diagnoses
        assert int(diagnoses[donor_id]["clinical_diagnosis"]) == int(source["cogdx"])
        assert int(diagnoses[donor_id]["pathological_diagnosis"]) == int(source["ad_adnc"])
        rin_supp = float(source["rin"])
        rin_author = float(confounds[donor_id]["rin"])
        pathology_group = "ADNC_HIGH" if source["ad_adnc"] == 1 else "ADNC_LOW"
        cogdx = int(source["cogdx"])
        donors.append(
            {
                "donor_id": donor_id,
                "pathology_group": pathology_group,
                "cogdx": cogdx,
                "clinical_group": {1: "NCI", 2: "MCI", 4: "AD_DEMENTIA"}.get(cogdx, "OTHER"),
                "protocol_rev": "Rev A" if donor_id in {2, 8, 12, 18} else "Rev D",
                "library_batch": int(confounds[donor_id]["libBatch"]),
                "sequencing_batch": "UNMAPPED_THREE_BATCHES",
                "study": str(source["study"]).strip(),
                "age_author_confounds": float(confounds[donor_id]["age"]),
                "rin_supplement": rin_supp,
                "rin_author_confounds": rin_author,
                "rin_conflict": rin_supp != rin_author,
            }
        )
    assert len(donors) == 21 and len({r["donor_id"] for r in donors}) == 21
    assert set(confounds) == set(diagnoses) == {r["donor_id"] for r in donors}
    assert Counter(r["pathology_group"] for r in donors) == {"ADNC_HIGH": 13, "ADNC_LOW": 8}
    assert Counter(r["clinical_group"] for r in donors) == {"AD_DEMENTIA": 10, "NCI": 6, "MCI": 5}

    def design_rank(subset: list[dict], group_name: str, *, protocol: bool, library_batch: bool,
                    age_rin: bool = False) -> tuple[int, int]:
        matrix = []
        for row in subset:
            values = [1, int(row[group_name] in {"ADNC_HIGH", "AD_DEMENTIA"})]
            if protocol:
                values.append(int(row["protocol_rev"] == "Rev A"))
            if library_batch:
                values.extend(int(row["library_batch"] == batch) for batch in range(2, 9))
            if age_rin:
                values.extend([row["age_author_confounds"], row["rin_author_confounds"]])
            matrix.append(values)
        array = np.asarray(matrix, dtype=float)
        return int(np.linalg.matrix_rank(array)), array.shape[1]

    clinical_strict = [r for r in donors if r["clinical_group"] in {"NCI", "AD_DEMENTIA"}]
    summary = {
        "n_unique_donors": len(donors),
        "pathology_by_protocol": {
            group: dict(Counter(r["protocol_rev"] for r in donors if r["pathology_group"] == group))
            for group in ("ADNC_HIGH", "ADNC_LOW")
        },
        "clinical_by_pathology": {
            group: dict(Counter(r["clinical_group"] for r in donors if r["pathology_group"] == group))
            for group in ("ADNC_HIGH", "ADNC_LOW")
        },
        "clinical_strict_by_protocol": {
            group: dict(Counter(r["protocol_rev"] for r in clinical_strict if r["clinical_group"] == group))
            for group in ("AD_DEMENTIA", "NCI")
        },
        "pathology_by_library_batch": {
            str(batch): dict(Counter(r["pathology_group"] for r in donors if r["library_batch"] == batch))
            for batch in range(1, 9)
        },
        "rank_intercept_pathology_protocol": design_rank(donors, "pathology_group", protocol=True, library_batch=False),
        "rank_intercept_clinical_protocol_strict": design_rank(clinical_strict, "clinical_group", protocol=True, library_batch=False),
        "rank_intercept_pathology_library_batch": design_rank(donors, "pathology_group", protocol=False, library_batch=True),
        "rank_intercept_pathology_protocol_library_batch": design_rank(donors, "pathology_group", protocol=True, library_batch=True),
        "rank_intercept_clinical_library_batch_strict": design_rank(clinical_strict, "clinical_group", protocol=False, library_batch=True),
        "rank_pathology_library_batch_age_rin": design_rank(donors, "pathology_group", protocol=False,
                                                             library_batch=True, age_rin=True),
        "rank_clinical_library_batch_age_rin_strict": design_rank(clinical_strict, "clinical_group",
                                                                   protocol=False, library_batch=True,
                                                                   age_rin=True),
        "mixed_pathology_library_batches": [batch for batch in range(1, 9) if
            len({r["pathology_group"] for r in donors if r["library_batch"] == batch}) == 2],
        "sequencing_batch_available": False,
        "library_batch_available": True,
        "rin_conflict_donors": [r["donor_id"] for r in donors if r["rin_conflict"]],
        "section_to_donor_available_in_table": False,
    }
    return donors, summary


def read_dharshini() -> tuple[list[dict], list[dict], dict]:
    workbook = load_workbook(DHARSHINI_TABLE, read_only=True, data_only=True)
    tissue = list(workbook["Tissue"].values)
    fields = tissue[0]
    sections = []
    for values in tissue[1:]:
        source = dict(zip(fields, values))
        if not isinstance(source["Donor ID"], int):
            continue
        sections.append(
            {
                "donor_id": int(source["Donor ID"]),
                "diagnosis": source["Disease Group"],
                "region": "BA9" if "BA9" in source["Brain Region"] else "BA17",
                "slide": source["Xenium Slide"],
                "rin": source["RIN"],
            }
        )
    assert len(sections) == 16
    assert len({r["donor_id"] for r in sections}) == 8
    assert all(Counter(r["diagnosis"] for r in sections if r["region"] == region) == {"AD": 4, "Control": 4}
               for region in ("BA9", "BA17"))
    assert all(Counter(r["diagnosis"] for r in sections if r["slide"] == slide and r["region"] == region)
               == {"AD": 1, "Control": 1}
               for slide in {r["slide"] for r in sections} for region in ("BA9", "BA17"))
    assert all({r["region"] for r in sections if r["donor_id"] == donor} == {"BA9", "BA17"}
               for donor in {r["donor_id"] for r in sections})

    with BASE_PANEL.open(newline="", encoding="utf-8-sig") as handle:
        base = {row["Genes"] for row in csv.DictReader(handle)}
    custom = {str(row[0]).strip() for row in list(workbook["Gene panel"].values)[1:] if row[0]}
    assert len(base) == 266 and len(custom) == 100 and base.isdisjoint(custom)
    panel = base | custom
    with gzip.open(G1, "rt", encoding="utf-8") as handle:
        g1 = {row["gene"] for row in csv.DictReader(handle, delimiter="\t")}
    assert len(g1) == 244
    with gzip.open(G5, "rt", encoding="utf-8") as handle:
        g5_rows = list(csv.DictReader(handle, delimiter="\t"))
    modules: dict[str, set[str]] = defaultdict(set)
    status = {}
    for row in g5_rows:
        module = row["module"]
        status[module] = (row["G5_primary_status"], row["G5_module_gate"])
        if row["one_to_one_mapping_status"].startswith("accepted") and row["human_one_to_one_symbol"]:
            modules[module].add(row["human_one_to_one_symbol"])
    overlap = [{"signature": "G1_signed_microglia", "frozen_genes": len(g1), "panel_genes": len(g1 & panel),
                "panel_fraction": round(len(g1 & panel) / len(g1), 6), "upstream_status": "FROZEN_G1",
                "genes_on_panel": ";".join(sorted(g1 & panel))}]
    for module in sorted(modules):
        genes = modules[module]
        overlap.append({"signature": module, "frozen_genes": len(genes), "panel_genes": len(genes & panel),
                        "panel_fraction": round(len(genes & panel) / len(genes), 6),
                        "upstream_status": "/".join(status[module]), "genes_on_panel": ";".join(sorted(genes & panel))})
    summary = {
        "sections": len(sections), "unique_donors": 8, "AD_donors": 4, "control_donors": 4,
        "regions": ["BA9", "BA17"], "slides": 4, "slide_region_diagnosis_balanced": True,
        "base_genes": len(base), "custom_genes": len(custom), "panel_union_genes": len(panel),
        "base_custom_overlap": len(base & custom),
        "panel_source_limitation": "Reconstructed from author custom 100 and official 10x base 266; not checked against actual Xenium gene_panel.json.",
    }
    return sections, overlap, summary


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    donors, avey = read_avey()
    sections, overlap, dharshini = read_dharshini()
    write_tsv(OUT / "avey_donor_design.tsv", donors,
              ["donor_id", "pathology_group", "cogdx", "clinical_group", "protocol_rev", "library_batch", "sequencing_batch", "study", "age_author_confounds", "rin_supplement", "rin_author_confounds", "rin_conflict"])
    write_tsv(OUT / "dharshini_donor_section.tsv", sections,
              ["donor_id", "diagnosis", "region", "slide", "rin"])
    write_tsv(OUT / "dharshini_frozen_panel_overlap.tsv", overlap,
              ["signature", "frozen_genes", "panel_genes", "panel_fraction", "upstream_status", "genes_on_panel"])
    inputs = [AVEY_TABLE, AVEY_CONFOUNDS, AVEY_DIAGNOSIS, DHARSHINI_TABLE, BASE_PANEL, G1, G5,
              RAW / "dharshini_zenodo_record.json"]
    manifest = [{"path": path.relative_to(ROOT).as_posix(), "bytes": path.stat().st_size,
                 "sha256": sha256(path)} for path in inputs]
    write_tsv(OUT / "input_sha256.tsv", manifest, ["path", "bytes", "sha256"])
    audit = {"audit_id": "G6-METADATA-2026-09-29-010", "effect_analysis_performed": False,
             "avey": avey, "dharshini": dharshini, "input_manifest": manifest}
    (OUT / "audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"avey": avey, "dharshini": dharshini, "panel_overlap": overlap}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
