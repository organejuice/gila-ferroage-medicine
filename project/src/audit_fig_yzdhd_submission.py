"""Preservation and submission audit; no statistical model is rerun."""
from pathlib import Path
import hashlib,json,re,zipfile
from urllib.parse import unquote
from docx import Document
from PIL import Image
from pypdf import PdfReader
import build_fig_yzdhd_submission as build
ROOT=build.ROOT; RAW=build.RAW; PROC=build.PROC; OUT=build.OUT
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    tests={}
    def check(name,value): tests[name]=bool(value)
    source_manifest=json.loads((PROC/'source_manifest.json').read_text(encoding='utf-8'))
    check('correct_project_sources_unchanged',all(sha(Path(s['source']))==s['sha256']==sha(ROOT/s['frozen_copy']) for s in source_manifest))
    meta=build.META; check('author_source_unchanged',sha(Path(meta['author_source']))==meta['author_source_sha256'])
    doc=Document(OUT/'manuscript/main_manuscript.docx')
    md=(OUT/'manuscript/main_manuscript.md').read_text(encoding='utf-8')
    text='\n'.join(p.text for p in doc.paragraphs)
    refs=json.loads((PROC/'final_references.json').read_text(encoding='utf-8'))
    check('reference_count_above_60',len(refs)>60)
    check('reference_dois_unique',len({r['doi'].lower() for r in refs})==len(refs))
    body,bibliography=md.split('## References\n',1)
    check('every_reference_cited',all(r['citation'] in body for r in refs))
    check('no_numeric_citation_markers',not re.search(r'\[\d+(?:[,–-]\d+)*\]',body))
    doi_urls=re.findall(r'https://doi.org/([^\s]+)',bibliography)
    check('doi_urls_preserve_encoded_identifiers',set(unquote(u).lower() for u in doi_urls)=={r['doi'].lower() for r in refs})
    ordered=sorted(refs,key=lambda r:(r['authors'][0].get('family','').casefold(),r['year'],r['canonical_title'].casefold()))
    check('alphabetical_bibliography',bibliography.strip().split('\n\n')==[build.format_ref(r) for r in ordered])
    abstract=body.split('## Abstract\n',1)[1].split('Keywords:',1)[0].strip()
    words=len(re.findall(r'\S+',abstract)); check('abstract_150_250_words',150<=words<=250)
    check('six_keywords',len(re.search(r'Keywords: ([^\n]+)',body)[1].split(';'))==6)
    check('all_nine_authors_present',all(a['full_name'] in text for a in meta['authors']))
    check('two_supplied_emails',meta['authors'][1]['email']=='author2@example.invalid' and meta['authors'][7]['email']=='author8@example.invalid')
    check('project_funding_and_conflicts_confirmed',meta['funding'] in text and meta['competing_interests'] in text)
    check('no_ADNI_scientific_transfer',not re.search(r'ADNI|AIBL|Fluid biomarker states|U19 AG024904|Clinical Dementia Rating',text))
    check('method_versions_unbroken','scTenifoldKnk 1.0.3 and scTenifoldNet 1.3' in text)
    check('matched_STRING_12_reference',any(r['id']==18 and r['doi']=='10.1093/nar/gkac1000' for r in refs))
    check('two_editable_main_tables',len(doc.tables)==2)
    from docx.oxml.ns import qn
    three_line=[]; white=[]
    for tbl in doc.tables:
        tb=tbl._tbl.tblPr.find(qn('w:tblBorders'))
        three_line.append(tb is not None and all(tb.find(qn('w:'+e)).get(qn('w:val'))==('single' if e in ['top','bottom'] else 'nil') for e in ['top','bottom','left','right','insideH','insideV']))
        for i,row in enumerate(tbl.rows):
            for c in row.cells:
                cb=c._tc.tcPr.find(qn('w:tcBorders'))
                three_line.append(cb is not None and all(cb.find(qn('w:'+e)).get(qn('w:val'))==('single' if (i==0 and e in ['top','bottom']) or (i==len(tbl.rows)-1 and e=='bottom') else 'nil') for e in ['top','bottom','left','right']))
                sh=c._tc.tcPr.find(qn('w:shd'))
                white.append(sh is None or sh.get(qn('w:fill')) in [None,'FFFFFF','auto'])
    check('main_tables_three_horizontal_rules_no_grid',all(three_line))
    check('main_table_headers_and_cells_no_colored_background',all(white))
    source=(RAW/'manuscript_en.md').read_text(encoding='utf-8').split('## References\n')[0]
    original_tables=[rows for kind,rows in build.blocks(source) if kind=='table']
    check('main_table_cell_values_preserved',[[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]==[[[build.plain(v) for v in row] for row in t] for t in original_tables])
    def norm(t,converted=False):
        if converted:
            for r in sorted(refs,key=lambda r:len(r['citation']),reverse=True): t=t.replace(r['citation'],'')
            t=re.sub(r'\([;\s]*\)','',t)
        else: t=re.sub(r'\[([\d,\s–-]+)\]','',t)
        t=build.plain(t)
        return re.sub(r'\s+([.,;])',r'\1',re.sub(r'\s+',' ',t)).strip()
    scientific=[]
    in_methods=False
    for kind,value in build.blocks(source):
        if kind=='h2' and value=='Materials and methods': in_methods=True
        if kind=='h2' and value=='Data availability': in_methods=False
        if in_methods and kind=='p': scientific.append(norm(value))
    import polish_fig_yzdhd_language
    approved_baseline=polish_fig_yzdhd_language.audited_baseline(md)
    normalized_body=norm(approved_baseline.split('## References\n',1)[0],True)
    language=json.loads((PROC/'language_style_audit.json').read_text(encoding='utf-8'))
    check('reviewed_language_transform_preserves_numbers_citations_bibliography',language['status']=='Passed')
    failures=[p[:95] for p in scientific if p not in normalized_body]
    check('inherited_scientific_prose_preserved_before_approved_language_edits',not failures)
    with zipfile.ZipFile(OUT/'supplementary/ESM_2.zip') as z:
        import csv,gzip,io
        files=[p for p in (RAW/'supplementary_tables').iterdir() if p.is_file()]
        check('all_eight_tables_byte_identical',all(hashlib.sha256(z.read('original_tables/'+p.name)).hexdigest()==sha(p) for p in files))
        csv_ok=[]
        for p in files:
            if not p.name.startswith('Table_S'): continue
            opener=gzip.open if p.suffix=='.gz' else open
            with opener(p,'rt',encoding='utf-8-sig',newline='') as f:
                original=list(csv.reader(f,delimiter=',' if p.suffix=='.csv' else '\t'))
            name=p.name.replace('.tsv.gz','.csv').replace('.tsv','.csv')
            converted=list(csv.reader(io.StringIO(z.read('tables/'+name).decode('utf-8-sig'))))
            csv_ok.append(original==converted)
        check('standard_CSV_preserves_all_source_fields',len(csv_ok)==8 and all(csv_ok))
        check('supplementary_table_zip_integrity',z.testzip() is None)
    with zipfile.ZipFile(OUT/'supplementary/ESM_3.zip') as z:
        manifest=json.loads(z.read('code_manifest.json'))
        check('code_zip_members_match_manifest',all(hashlib.sha256(z.read('project/'+p['path'])).hexdigest()==p['sha256'] for p in manifest))
        check('code_source_files_unchanged',all(sha(ROOT/p['path'])==p['sha256'] for p in manifest))
        check('no_raw_molecular_data_in_code_zip',not any(n.endswith(('.h5','.h5ad','.rds','.RData','.fastq','.fastq.gz')) for n in z.namelist()))
        check('code_zip_integrity',z.testzip() is None)
    images=[]
    for i in range(1,5):
        im=Image.open(OUT/f'figures/Fig{i}.tif'); dpi=im.info['dpi']; width=im.width/dpi[0]*25.4
        images.append(im.mode=='RGB' and min(dpi)>=599 and abs(width-174)<.1)
    check('four_submission_figures_600dpi_RGB_174mm',all(images))
    check('four_figures_embedded_in_Word',len(doc.inline_shapes)==4)
    import csv, math
    data=OUT/'figures/source_data'
    with (data/'Fig4a_estimable_reproduction_rates.tsv').open(encoding='utf-8',newline='') as f: a=list(csv.DictReader(f,delimiter='\t'))
    with (data/'Fig4b_estimable_storage_comparisons.tsv').open(encoding='utf-8',newline='') as f: b=list(csv.DictReader(f,delimiter='\t'))
    fields=['lodo_weighted_median_unconditional_rate','bootstrap_weighted_median_unconditional_rate']
    check('Fig4_display_data_finite_and_expected_counts',len(a)==20 and len(b)==42 and all(math.isfinite(float(r[k])) for r in a for k in fields) and all(math.isfinite(float(r['reference_tail_fraction'])) for r in b))
    exclusions=json.loads((data/'Fig4_exclusion_log.json').read_text())
    check('Fig4_nonestimable_exclusions_documented',exclusions['panel_a']['before']==29 and len(exclusions['panel_a']['excluded'])==9 and exclusions['panel_b']['before']==58 and len(exclusions['panel_b']['excluded'])==16)
    with (ROOT/'results/2026-10-11-virtual-ko-scientific-review-035/target_verification_conclusions.tsv').open(encoding='utf-8',newline='') as f: original_a=list(csv.DictReader(f,delimiter='\t'))
    with (ROOT/'results/2026-10-11-virtual-ko-matching-correction-036/functional_response_comparison.tsv').open(encoding='utf-8',newline='') as f: original_b=list(csv.DictReader(f,delimiter='\t'))
    def finite(v):
        try: return math.isfinite(float(v))
        except ValueError: return False
    with (data/'Fig2c_astrocyte.tsv').open(encoding='utf-8',newline='') as f: c=list(csv.DictReader(f,delimiter='\t'))
    with (ROOT/'results/2026-10-13-manuscript/supplementary_tables/Table_S4_astrocyte_modules.tsv').open(encoding='utf-8',newline='') as f: original_c=list(csv.DictReader(f,delimiter='\t'))
    cf=['logFC','CI_low','CI_high','FDR_BH_7']
    expected_c={r['module']:r for r in original_c if all(finite(r[k]) for k in cf)}
    check('Fig2c_all_six_estimable_effects_CIs_original_FDR_preserved',len(c)==6 and {r['module'] for r in c}==set(expected_c) and all(finite(r[k]) and abs(float(r[k])-float(expected_c[r['module']][k]))<1e-14 for r in c for k in cf))
    key=lambda r:tuple(r[k] for k in ['analysis','context','target'])
    expected_a={key(r):r for r in original_a if all(finite(r[k]) for k in fields)}
    expected_b={key(r)+(r['seed'],):r for r in original_b if r['variant']=='corrected_full_CP10K' and 'Storage' in r['function'] and finite(r['reference_tail_fraction'])}
    check('Fig4_all_finite_values_preserved_including_zero',set(map(key,a))==set(expected_a) and {key(r)+(r['seed'],) for r in b}==set(expected_b) and all(abs(float(r[k])-float(expected_a[key(r)][k]))<1e-14 for r in a for k in fields) and all(abs(float(r['reference_tail_fraction'])-float(expected_b[key(r)+(r['seed'],)]['reference_tail_fraction']))<1e-14 for r in b))
    check('four_figures_cited_in_order',[body.index('See Fig. '+str(i)+'.') for i in range(1,5)]==sorted(body.index('See Fig. '+str(i)+'.') for i in range(1,5)))
    check('three_online_resources_cited',all('Online Resource '+str(i) in body for i in range(1,4)))
    check('statements_and_AI_disclosure','Statements and Declarations' in text and 'OpenAI Codex assisted' in text)
    check('clinical_trial_exact_not_applicable_statement',text.count('Clinical trial number: not applicable.')==1 and md.count('Clinical trial number: not applicable.')==1 and 'trialregister.nl' not in md)
    check('all_eight_supplementary_table_captions_in_main_manuscript',all(c in text for c in build.SUPPLEMENTARY_TABLE_CAPTIONS))
    with zipfile.ZipFile(OUT/'supplementary/ESM_2.zip') as z:
        index=z.read('TABLE_CAPTIONS.md').decode('utf-8')
        check('supplementary_caption_index_matches_manuscript',all(c in index for c in build.SUPPLEMENTARY_TABLE_CAPTIONS))
    check('no_placeholder_text',not re.search(r'TODO|TBD|AUTHOR_INPUT_NEEDED|INSERT HERE',text))
    counts={}
    for stem in ['main_manuscript','title_page','cover_letter','ESM_1']:
        folder=ROOT/'tmp/FIG_YZDHD_20261014'/('render_'+stem)
        pdf=folder/(stem+'.pdf'); count=len(PdfReader(pdf).pages); counts[stem]=count
        check('complete_render_'+stem,len(list(folder.glob('page-*.png')))==count and count>0)
    value={'status':'Passed' if all(tests.values()) else 'Failed','checks':tests,'reference_count':len(refs),'abstract_words':words,'inherited_scientific_paragraphs_checked':len(scientific),'prose_preservation_failures':failures,'rendered_pages':counts,'boundary':'Preservation and editorial audits, no patient-level or molecular model rerun'}
    (PROC/'package_audit.json').write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'status':value['status'],'checks':len(tests),'failed':[k for k,v in tests.items() if not v],'abstract_words':words,'pages':counts},ensure_ascii=False))
    if not all(tests.values()): raise SystemExit(1)
if __name__=='__main__': main()
