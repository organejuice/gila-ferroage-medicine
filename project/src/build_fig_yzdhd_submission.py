"""Build the correct YZDHD submission from frozen text and verified references."""
from pathlib import Path
import csv, gzip, hashlib, html, json, re, shutil, zipfile
from urllib.parse import quote
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
RAW=ROOT/'data/raw/FIG_YZDHD_20261014'
PROC=ROOT/'data/processed/FIG_YZDHD_20261014'
OUT=ROOT/'results/2026-10-14-FIG-YZDHD'
QA=ROOT/'tmp/FIG_YZDHD_20261014'
META=json.loads((PROC/'author_metadata.json').read_text(encoding='utf-8'))
TITLE=META['title']; JOURNAL=META['journal']; NAMES=[a['full_name'] for a in META['authors']]
CONTACT='Corresponding author: Rui Qie; qierui@hljucm.edu.cn'
SUPPLEMENTARY_TABLE_CAPTIONS=[
    'Table S1. Human microglial Alzheimer’s disease association signature.',
    'Table S2. Senescence-associated and inflammatory transcriptional module responses in amyloid-β-exposed induced pluripotent stem cell-derived microglia.',
    'Table S3. Iron-handling and lipid-peroxidation-associated transcriptional module responses in induced pluripotent stem cell-derived microglia.',
    'Table S4. Disease-associated transcriptional module effects in APP/PS1-21 mouse astrocytes.',
    'Table S5. Accepted experimental-record compound–target pairs.',
    'Table S6. Structural assessment and molecular docking results.',
    'Table S7. Virtual knockout robustness under donor omission and bootstrap resampling.',
    'Table S8. Corrected matched-reference comparisons for virtual knockout responses.'
]

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def plain(t):
    t=html.unescape(t)
    return re.sub(r'\[([^\]]+)\]\((https?://[^\s)]+)\)',r'\1 (\2)',t)
def make_doc(size=11,spacing=1.25):
    d=Document(); s=d.sections[0]
    s.page_width=Inches(8.27); s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.75)
    s.left_margin=s.right_margin=Inches(.82)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        st=d.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(size if name=='Normal' else 14 if name=='Title' else 12)
        st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=spacing
        st.paragraph_format.widow_control=True
        if name.startswith('Heading'): st.paragraph_format.keep_with_next=True; st.font.bold=True
        if name=='Title': st.font.bold=True; st.paragraph_format.line_spacing=1.05
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'SimSun')
    for b in list(d.styles.element.iter(qn('w:pBdr'))): b.getparent().remove(b)
    p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE')
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text='1'; r.append(t); f.append(r); p._p.append(f)
    d.core_properties.author='; '.join(NAMES); d.core_properties.title=TITLE; d.core_properties.subject=JOURNAL
    return d
def paragraph(d,text,style=None):
    p=d.add_paragraph(plain(text),style=style)
    p.paragraph_format.keep_together=False
    return p
def table(d,rows,widths):
    t=d.add_table(rows=1,cols=len(rows[0])); t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge)
        el.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil')
        el.set(qn('w:sz'),'12' if edge in ['top','bottom'] else '0')
        el.set(qn('w:color'),'000000'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]); cell.text=plain(value)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge)
                outer=(i==0 and edge=='top') or (i==len(rows)-1 and edge=='bottom')
                header=i==0 and edge=='bottom'
                el.set(qn('w:val'),'single' if outer or header else 'nil')
                el.set(qn('w:sz'),'12' if outer else '6' if header else '0')
                el.set(qn('w:color'),'000000'); borders.append(el)
            pr.append(borders)
            sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),'FFFFFF'); pr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1
                p.paragraph_format.keep_with_next=i<len(rows)-1
                for r in p.runs: r.font.size=Pt(9.5); r.bold=i==0
        pr=t.rows[i]._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit'); pr.append(ns)
        if i==0: pr.append(OxmlElement('w:tblHeader'))
    return t
def blocks(text):
    lines=text.splitlines(); i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x) for x in row): rows.append(row)
                i+=1
            yield 'table',rows; continue
        if line.startswith('#'):
            n=len(line)-len(line.lstrip('#')); yield 'h'+str(n),line[n:].strip(); i+=1; continue
        if line.startswith('- '): yield 'bullet',line[2:]; i+=1; continue
        lineset=[line]; i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','- ')):
            lineset.append(lines[i].strip()); i+=1
        yield 'p',' '.join(lineset)
def initial(given): return ''.join(v[0].upper() for v in re.findall(r'[A-Za-zÀ-ž]+',given or ''))
def refs_and_plans():
    orig=json.loads((PROC/'original_canonical_references.json').read_text(encoding='utf-8'))
    new=json.loads((PROC/'new_reference_candidates.json').read_text(encoding='utf-8'))
    refs=(orig if isinstance(orig,list) else orig['references'])+(new if isinstance(new,list) else new['references'])
    assert len(refs)>60
    assert len({r['doi'].lower() for r in refs})==len(refs)
    for r in refs:
        r['canonical_title']=r.get('canonical_title',r.get('title',''))
        r['page']=r.get('page') or r.get('pages') or r.get('article_number') or ''
        authors=r['authors']; n=len(authors)
        r['citation']=(authors[0].get('family',authors[0].get('collective','')) if n==1 else authors[0]['family']+' and '+authors[1]['family'] if n==2 else authors[0]['family']+' et al.')+' '+str(r['year'])
    groups={}
    for r in refs: groups.setdefault(r['citation'],[]).append(r)
    for same in groups.values():
        if len(same)>1:
            for j,r in enumerate(sorted(same,key=lambda v:v['canonical_title'])):
                r['suffix']=chr(97+j); r['citation']+=r['suffix']
    plans=json.loads((PROC/'new_reference_insertion_plans.json').read_text(encoding='utf-8'))
    return refs,plans
def format_ref(r):
    authors=[]
    for a in r['authors'][:6]:
        if a.get('collective'): authors.append(a['collective']); continue
        authors.append(a['family']+' '+initial(a.get('given',''))+(' '+a['suffix'] if a.get('suffix') else ''))
    if len(r['authors'])>6: authors.append('et al')
    loc=str(r.get('volume') or '')+(':'+str(r.get('page') or '') if r.get('page') else '')
    return ', '.join(authors)+f" ({r['year']}{r.get('suffix','')}) "+html.unescape(r.get('display_title',r['canonical_title']))+'. '+r['journal']+(' '+loc if loc else '')+'. https://doi.org/'+quote(r['doi'],safe='/()')
def replace_cites(text,refmap):
    def sub(m):
        ids=[]
        for token in m[1].replace('–','-').split(','):
            p=token.strip().split('-'); ids+=list(range(int(p[0]),int(p[-1])+1))
        return '('+'; '.join(sorted({refmap[i]['citation'] for i in ids},key=str.casefold))+')'
    return re.sub(r'\[([\d,\s–-]+)\]',sub,text)
def main_text(refs,plans):
    source=(RAW/'manuscript_en.md').read_text(encoding='utf-8')
    body=source.split('## References\n')[0]
    body=body.replace(body.splitlines()[0],'# '+TITLE,1)
    body=re.sub(r'Keywords: .*', 'Keywords: Alzheimer disease; glial transcriptomics; histone lactylation; network pharmacology; virtual knockout; Yizhi Dihuang Decoction',body,count=1)
    abstract=body.split('## Abstract\n',1)[1].split('Keywords:',1)[0]
    new_abstract=abstract.replace('H4K12 lactylation','histone H4 lysine 12 lactylation').replace('In APP/PS1-21 astrocytes','In amyloid-model mouse astrocytes').replace('frozen ADMET prediction rules','frozen absorption, distribution, metabolism, excretion and toxicity (ADMET) prediction rules')
    body=body.replace(abstract,new_abstract,1)
    changes=[]
    for plan in plans:
        anchor=plan.get('anchor',plan.get('paragraph_anchor',''))
        addition=plan.get('text',plan.get('addition',''))
        ids=plan.get('ref_ids',plan.get('reference_ids',[]))
        if ids and not re.search(r'\[\d',addition): addition=addition.rstrip('.')+' ['+','.join(map(str,ids))+'].'
        paragraphs=body.split('\n\n'); matches=[i for i,p in enumerate(paragraphs) if anchor and anchor in p]
        assert len(matches)==1,('Insertion anchor must match one paragraph',anchor,matches)
        if plan.get('action')=='cite_sentence':
            p=paragraphs[matches[0]]; start=p.index(anchor)
            end=start+len(anchor)-1 if anchor.endswith('.') else start+re.search(r'\.(?=\s|$)',p[start:]).start()
            p=p[:end]+' ['+','.join(map(str,ids))+']'+p[end:]
            paragraphs[matches[0]]=p
        else: paragraphs[matches[0]]+=' '+addition
        body='\n\n'.join(paragraphs); changes.append(plan)
    # Keep figures as separate publication files, cite their role without new findings.
    fig_anchors=[('The analysis linked human AD microglial expression','Fig. 1'),('In the iMGL time course, iron-uptake scores increased','Fig. 2'),('Within the human 244-gene subset, the E1 layer contained','Fig. 3'),('Final scientific review examined fifteen targets','Fig. 4')]
    for anchor,label in fig_anchors:
        paras=body.split('\n\n'); index=next(i for i,p in enumerate(paras) if p.startswith(anchor)); paras[index]+=' See '+label+'.'; body='\n\n'.join(paras)
    body=body.replace('Derived tables, source-specific retrieval status, frozen gene sets, input hashes, and exclusion records are retained in the project archive. Public deposition of that archive has not yet been established.','Derived tables are supplied in Online Resource 2, with source hashes and interpretation notes in Online Resource 1. Analysis code and configuration are provided in Online Resource 3. Public repository deposition has not yet been established.')
    body=body.replace('The archive supports local traceable reruns; a public code release and verified DVC remote restore remain to be completed.','The code archive in Online Resource 3 records the local workflow and its dependencies; it does not include raw molecular count matrices. Public deposition and cross-machine restoration of the complete upstream research archive have not been established.')
    disclosure='## Use of artificial intelligence\n\nOpenAI Codex assisted with manuscript drafting, language editing, bibliographic searches, citation checks and preparation of submission files. It was not used to generate new biological observations in this journal-conversion step. The authors are responsible for evaluating the final text, references and interpretations.\n\n'
    declarations='## Statements and Declarations\n\n### Funding\n\n'+META['funding']+'\n\n### Competing interests\n\n'+META['competing_interests']+'\n\n### Ethics approval and consent to participate\n\nThis study reanalysed public molecular datasets and published experimental records. It involved no new human recruitment or animal experimentation. Ethics and consent requirements for the original data collection are described in the source studies. No identifiable individual-level clinical information is reported.\n\n### Consent for publication\n\nNot applicable. The manuscript contains no identifiable individual data.\n\n### Author contributions\n\n'+META['contributions']+'\n\n'
    body=body.replace('## Results\n',disclosure+'## Results\n',1)+declarations
    rmap={r['id']:r for r in refs}; cited=set()
    for m in re.finditer(r'\[([\d,\s–-]+)\]',body):
        for token in m[1].replace('–','-').split(','):
            ends=token.strip().split('-'); cited.update(range(int(ends[0]),int(ends[-1])+1))
    assert cited==set(rmap),('Uncited references',set(rmap)-cited)
    body=re.sub(r'\[([\d,\s–-]+)\]\s+\[([\d,\s–-]+)\]',lambda m:'['+m[1]+','+m[2]+']',body)
    text=replace_cites(body,rmap)+'## References\n\n'+'\n\n'.join(format_ref(r) for r in sorted(refs,key=lambda r:(r['authors'][0].get('family','').casefold(),r['year'],r['canonical_title'].casefold())))+'\n'
    (PROC/'final_references.json').write_text(json.dumps(refs,ensure_ascii=False,indent=2),encoding='utf-8')
    (PROC/'editorial_insertions.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
    import polish_fig_yzdhd_language
    return polish_fig_yzdhd_language.apply(text)
def manuscript(text):
    d=make_doc()
    for kind,value in blocks(text):
        if kind=='h1':
            paragraph(d,value,'Title'); paragraph(d,', '.join(NAMES[:-1])+', Rui Qie*'); paragraph(d,META['authors'][0]['affiliation']); paragraph(d,META['authors'][-1]['affiliation']); paragraph(d,CONTACT)
        elif kind.startswith('h'): d.add_heading(value,min(3,int(kind[1])-1))
        elif kind=='table': table(d,value,[1.03,1.12,.68,.74,1.23,1.78] if len(value[0])==6 else [1.62,.82,.95,1.81,1.38])
        else:
            p=paragraph(d,value,'List Bullet' if kind=='bullet' else None)
            if value.startswith('Table '): p.paragraph_format.keep_with_next=True; p.runs[0].bold=True
            if kind=='p' and 'https://doi.org/' in value: p.paragraph_format.space_after=Pt(4)
    captions_path=OUT/'figures/figure_captions.json'
    assert captions_path.exists(),'Figure captions must be finalized before main DOCX'
    captions=json.loads(captions_path.read_text(encoding='utf-8'))
    d.add_page_break()
    d.add_heading('Figures and legends',1)
    for item in (captions if isinstance(captions,list) else captions['figures']):
        if item!=captions[0]: d.add_page_break()
        image_paragraph=d.add_paragraph()
        image_paragraph.add_run().add_picture(str(OUT/'figures'/Path(item['pdf']).with_suffix('.png')),width=Inches(6.6))
        image_paragraph.paragraph_format.space_after=Pt(4)
        image_paragraph.paragraph_format.keep_with_next=True
        legend=item.get('caption',item.get('legend','')).rstrip('.')
        p=d.add_paragraph(); prefix=re.match(r'Fig\. \d+',legend)[0]
        p.add_run(prefix).bold=True; p.add_run(legend[len(prefix):])
        p.paragraph_format.line_spacing=1; p.paragraph_format.space_after=Pt(4)
        for r in p.runs: r.font.size=Pt(9.5)
    d.add_heading('Supplementary material captions',1).paragraph_format.page_break_before=True
    for caption in ['Online Resource 1 Supplementary table guide and interpretation boundaries','Online Resource 2 Eight derived supplementary tables in CSV format, source-identical original tables, estimable Fig. 2c and Fig. 4 plot data and exclusion records','Online Resource 3 Recorded analysis code, tests, configuration, environment files and source-to-claim provenance']:
        paragraph(d,caption)
    d.add_heading('Supplementary table captions',2)
    for caption in SUPPLEMENTARY_TABLE_CAPTIONS:
        paragraph(d,caption)
    (OUT/'manuscript').mkdir(exist_ok=True); d.save(OUT/'manuscript/main_manuscript.docx')
    (OUT/'manuscript/main_manuscript.md').write_text(text,encoding='utf-8')
def administrative():
    folder=OUT/'administrative'; folder.mkdir(exist_ok=True)
    d=make_doc(11,1.1); paragraph(d,TITLE,'Title'); paragraph(d,', '.join(NAMES))
    paragraph(d,'Authors 1–8: '+META['authors'][0]['affiliation']); paragraph(d,'Author 9: '+META['authors'][-1]['affiliation']); paragraph(d,CONTACT)
    for title,value in [('Author contributions',META['contributions']),('Funding',META['funding']),('Competing interests',META['competing_interests']),('Acknowledgements','We thank the investigators who generated the public molecular datasets and the maintainers of the referenced open research resources.')]:
        d.add_heading(title,1); paragraph(d,value)
    d.save(folder/'title_page.docx')
    d=make_doc(11,1.1); paragraph(d,'14 October 2026'); paragraph(d,'Dear Editors,')
    paragraph(d,'Please consider our manuscript “'+TITLE+'” as a Research article for '+JOURNAL+'.')
    paragraph(d,'The study integrates donor-level human glial transcriptomics with model-specific epigenetic and transcriptional programs, compound–target provenance, structural assessments and virtual gene knockouts. Three human cohorts provide five region units and a directionally stable 244-gene association subset. Complementary model analyses define distinct glial response programs that connect lactylation, senescence and iron homeostasis to Yizhi Dihuang Decoction-related pharmacological hypotheses.')
    paragraph(d,'The integrative analysis also identifies limits that matter for functional interpretation: the primary candidate network does not show significant proximity to the human signature, and knockout responses remain dependent on donor sampling, expression support and matched-reference behavior. Together, these findings delineate context-dependent glial responses and candidate mechanisms that require direct measurements of formula exposure, target engagement and biological effects.')
    paragraph(d,'The manuscript addresses integrative transcriptomic analysis and functional hypothesis evaluation in Alzheimer’s disease. Supplementary tables, source hashes and the recorded analysis code support assessment of the results. We recognize that the current candidate hypotheses require subsequent experimental validation.')
    paragraph(d,'The authors report no funding and no competing interests. Author information and contribution roles are supplied in the title page and submission metadata. Data availability, ethical context and AI-assisted manuscript preparation are disclosed in the manuscript.')
    paragraph(d,'Sincerely,'); paragraph(d,'Rui Qie\nOn behalf of the authors\nqierui@hljucm.edu.cn'); d.save(folder/'cover_letter.docx')
    with (folder/'author_submission_metadata.tsv').open('w',encoding='utf-8',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(META['authors'][0]),delimiter='\t'); writer.writeheader(); writer.writerows(META['authors'])
def supplementary():
    folder=OUT/'supplementary'; folder.mkdir(exist_ok=True); QA.mkdir(parents=True,exist_ok=True)
    d=make_doc(11,1.2); paragraph(d,'Online Resource 1','Title'); paragraph(d,TITLE); paragraph(d,JOURNAL); paragraph(d,'; '.join(NAMES)); paragraph(d,META['authors'][-1]['affiliation']); paragraph(d,CONTACT)
    sections=[('Supplementary material organization','Online Resource 2 contains eight source-identical machine-readable tables, their provenance manifest and a metadata file. Online Resource 3 provides analysis scripts, test modules, configuration, environment files and the frozen source-to-claim map. Source molecular count matrices and external database exports are not redistributed in the code archive.'),('Table S1 Human association signature','244 genes with effects estimable in all five region units and direction retained after each whole cohort was omitted. Sign stability does not require significant results in every reduced analysis.'),('Table S2 Microglial senescence time course','iMGL sample-level SASP, inflammatory and operational core comparisons. Culture replicates are not independent human donors; standardized gene scores are transcriptional proxies.'),('Table S3 Microglial iron and lipid programs','Iron uptake, storage, export, ferritinophagy, LPO-associated and ferro-aging reference modules. Module-mean scores were standardized across samples. RNA does not measure iron concentration, lipid flux or cell death.'),('Table S4 Astrocyte primary modules','APP/PS1-21 astrocyte disease coefficients and confidence intervals. The fixed BH family includes seven modules, including the nonestimable storage module. Historical storage/ferritinophagy modules did not enter subsequent candidate selection.'),('Table S5 Compound and target pairs','Eighty unique accepted experimental-record connections across twenty program-connected targets. Pair counts differ from assay/source records; E1 does not establish direct cellular binding, formula exposure or therapeutic direction.'),('Table S6 Structural scope','Target-specific structural preparation, candidate docking calculations and reference-pose checks. Historical platform scores and local docking are exploratory. Reference RMSD concerns reference ligands, and scores must not be compared as affinities across targets.'),('Table S7 Virtual knockout robustness','Fifteen targets across twenty-nine combinations assessed with donor omission and bootstrap analyses. Planned and estimable denominators are distinguished; absent target, low expression and no outgoing edges are not negative effects.'),('Table S8 Corrected reference comparisons','Full-library CP10K matching and finite network expansion for low-expression targets. Descriptive reference tail fractions are not formal specificity P values. The historical two-gene storage comparison does not validate an entire ferro-aging pathway. The correction did not rerun the preceding full donor-resampling and enrichment analyses.'),('Reproducibility and interpretation','Supplementary tables preserve the original numerical results. Figures summarize these results without additional model fitting. Code availability is distinct from a self-contained data archive: raw data acquisition, source exports and environment-specific paths remain required for upstream reruns. The supplied code and configuration support reconstruction of the manuscript outputs; the complete upstream biological analysis has not been independently reproduced on another machine.')]
    for head,text in sections: d.add_heading(head,1); paragraph(d,text)
    d.save(QA/'ESM_1.docx')
    metadata={'title':TITLE,'journal':JOURNAL,'authors':NAMES,'corresponding_author':CONTACT,'corresponding_affiliation':META['authors'][-1]['affiliation'],'content':'Eight frozen derived supplementary tables in standard CSV plus source-identical original formats','source_table_manifest':'original_tables/table_provenance.json'}
    with zipfile.ZipFile(folder/'ESM_2.zip','w',zipfile.ZIP_DEFLATED) as z:
        import io
        for p in sorted((RAW/'supplementary_tables').glob('*')):
            z.write(p,'original_tables/'+p.name)
            if p.name.startswith('Table_S'):
                opener=gzip.open if p.suffix=='.gz' else open
                with opener(p,'rt',encoding='utf-8-sig',newline='') as f:
                    rows=list(csv.reader(f,delimiter=',' if p.suffix=='.csv' else '\t'))
                csv_buffer=io.StringIO(newline=''); writer=csv.writer(csv_buffer); writer.writerows(rows)
                name=p.name.replace('.tsv.gz','.csv').replace('.tsv','.csv')
                z.writestr('tables/'+name,csv_buffer.getvalue().encode('utf-8-sig'))
        z.writestr('TABLE_CAPTIONS.md','# Supplementary table captions\n\n'+'\n\n'.join(SUPPLEMENTARY_TABLE_CAPTIONS)+'\n')
        z.writestr('SUBMISSION_METADATA.json',json.dumps(metadata,ensure_ascii=False,indent=2))
        for p in sorted([*(OUT/'figures/source_data').glob('Fig4*'),*(OUT/'figures/source_data').glob('Fig2c*')]):
            z.write(p,'figure_data/'+p.name)
        z.writestr('README.md','# Supplementary tables\n\nUse tables/*.csv for standard spreadsheet access. figure_data/Fig2c_astrocyte.tsv contains the six estimable modules displayed in Fig. 2c; the original seven-test BH correction is retained, with exclusion details in Fig2c_exclusion_log.json. figure_data/Fig4a_estimable_reproduction_rates.tsv and Fig4b_estimable_storage_comparisons.tsv contain only the finite statistics displayed in Fig. 4; exclusions and reasons are recorded in Fig4_exclusion_log.json. The two panels use different eligible sets. original_tables/ retains byte-identical source CSV/TSV/gzip tables and their provenance, including nonestimable assessment records. CSV conversion preserves every data field as text and does not round numerical strings. SUBMISSION_METADATA.json identifies the article, journal, authors and correspondence.\n')
    included=[]
    for directory in ['src','tests','configs']:
        included+=sorted(p for p in (ROOT/directory).rglob('*') if p.is_file() and '__pycache__' not in p.parts and p.suffix not in ['.pyc','.png','.pdf','.zip','.xlsx'])
    included += [ROOT/p for p in ['README.md','params.yaml','environment.yml','requirements.txt','dvc.yaml','dvc.lock'] if (ROOT/p).exists()]
    metadata['content']='Frozen project code and configurations; raw counts and upstream result archives not included'
    with zipfile.ZipFile(folder/'ESM_3.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in included: z.write(p,'project/'+p.relative_to(ROOT).as_posix())
        z.write(RAW/'claim_evidence.tsv','claim_evidence.tsv')
        z.write(RAW/'provenance.json','original_manuscript_provenance.json')
        z.writestr('SUBMISSION_METADATA.json',json.dumps(metadata,ensure_ascii=False,indent=2))
        z.writestr('code_manifest.json',json.dumps([{'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p)} for p in included],ensure_ascii=False,indent=2))
        z.writestr('REBUILD_NOTES.md','# Rebuild boundaries\n\nThe code archive contains local project implementations and configuration, but not public raw count matrices or full frozen intermediate outputs. Obtain source datasets and database exports using the registered accessions, preserve their hashes, and resolve platform-specific paths before upstream reruns. Run src/build_fig_yzdhd_submission.py only within the full editorial workspace with its frozen inputs and verified bibliography. No self-contained cross-machine replication is claimed.\n')
    (PROC/'code_archive_manifest.json').write_text(json.dumps([{'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p)} for p in included],ensure_ascii=False,indent=2),encoding='utf-8')
def main():
    OUT.mkdir(exist_ok=True)
    refs,plans=refs_and_plans(); text=main_text(refs,plans)
    manuscript(text); administrative(); supplementary()
    print(json.dumps({'references':len(refs),'authors':len(NAMES),'main_tables':2,'supplementary_tables':8,'scientific_source':'YZDHD project','ADNI_content_used':'author metadata only'},ensure_ascii=False))
if __name__=='__main__': main()
